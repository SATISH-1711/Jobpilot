import os

from docx import Document
from pypdf import PdfReader

MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MIN_TEXT_LENGTH = 50


class ResumeParseError(Exception):
    """Raised with a message that is safe to show to the user."""


def extract_text(uploaded_file):
    extension = os.path.splitext(uploaded_file.name)[1].lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise ResumeParseError("Unsupported file type. Upload a PDF, DOCX or TXT file.")
    if uploaded_file.size > MAX_FILE_SIZE:
        raise ResumeParseError("File is too large. The limit is 5 MB.")

    try:
        if extension == ".pdf":
            reader = PdfReader(uploaded_file)
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        elif extension == ".docx":
            document = Document(uploaded_file)
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        parts.append(cell.text)
            text = "\n".join(parts)
        else:
            text = uploaded_file.read().decode("utf-8", errors="ignore")
    except Exception:
        raise ResumeParseError(
            "Could not read this file. It may be corrupted or password protected."
        )

    text = text.replace("\x00", "").strip()
    if len(text) < MIN_TEXT_LENGTH:
        raise ResumeParseError(
            "Almost no text was found. Scanned or image-only PDFs are not supported."
        )
    return text